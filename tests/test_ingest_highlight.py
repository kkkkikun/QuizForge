"""loader 高亮提取测试（黄色高亮 = 答案标记）。"""

import os
import tempfile

from quizforge.ingest import load_docx


def _make_docx_with_highlight(answer_letter: str) -> str:
    """造一个 CET 风格题：听力模块 + 1 题四选项，answer_letter 选项整段黄色高亮。"""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    doc.add_paragraph("听力")
    doc.add_paragraph("Unit 1")
    doc.add_paragraph("News report 1")
    doc.add_paragraph("1、")
    for letter in "ABCD":
        p = doc.add_paragraph(f"{letter}、选项{letter}")
        if letter == answer_letter:
            for r in p.runs:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def test_loader_captures_highlight():
    path = _make_docx_with_highlight("C")
    try:
        blocks = load_docx(path)
    finally:
        os.remove(path)
    # 找到四选项块
    opt_blocks = [b for b in blocks if b.text.startswith(("A、", "B、", "C、", "D、"))]
    assert len(opt_blocks) == 4
    # 只有 C 选项块带高亮
    hl = [b for b in opt_blocks if b.highlight]
    assert len(hl) == 1
    assert hl[0].text.startswith("C、")
    assert "选项C" in hl[0].highlight


def test_loader_no_highlight_default_empty():
    path = _make_docx_with_highlight("A")
    try:
        blocks = load_docx(path)
    finally:
        os.remove(path)
    # 非高亮选项的 highlight 为空
    for b in blocks:
        if b.text.startswith("B、"):
            assert b.highlight == ""
            return
    assert False, "未找到 B 选项块"
