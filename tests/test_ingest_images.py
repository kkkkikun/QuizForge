"""图片抽取：docx 段落内图片 → base64 入 Block.images（题目配图不再丢）。"""

import io

import pytest
from docx import Document
from docx.shared import Pt

from quizforge.ingest import load_docx

# 1×1 红色 PNG
_PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
    "53de0000000c4944415408d763f8cfc000000301010018dd8db00000000049"
    "454e44ae426082"
)


def _make_docx_with_image(path, wrap_pict=False):
    """构造带一张图片的 docx（DrawingML 或 VML 包裹）。"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_paragraph("题干：如下图所示")
    run = doc.add_paragraph().add_run()
    if wrap_pict:
        # VML v:imagedata 引用（老式；python-docx 不内置 VML 前缀，用 parse_xml 构造）
        from docx.oxml import parse_xml
        rId, _ = doc.part.get_or_add_image(io.BytesIO(_PNG_1PX))
        nsdecl = (
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:v="urn:schemas-microsoft-com:vml"'
        )
        pict = parse_xml(f'<w:pict {nsdecl}><v:shape><v:imagedata r:id="{rId}"/></v:shape></w:pict>')
        run._element.append(pict)
    else:
        # DrawingML inline 图片（python-docx 标准路径）
        run.add_picture(io.BytesIO(_PNG_1PX))
    doc.add_paragraph("A、选项甲")
    doc.add_paragraph("B、选项乙")
    doc.save(path)


def test_docx_image_extracted_as_base64(tmp_path):
    """DrawingML 图片 → image Block 携带 base64 数据。"""
    p = tmp_path / "img.docx"
    _make_docx_with_image(str(p))
    blocks = load_docx(str(p))
    imgs = [b for b in blocks if b.kind == "image"]
    assert len(imgs) == 1
    import base64
    assert base64.b64decode(imgs[0].images[0]) == _PNG_1PX


def test_docx_vml_image_extracted(tmp_path):
    """VML v:imagedata 引用的图片同样提取。"""
    p = tmp_path / "vml.docx"
    _make_docx_with_image(str(p), wrap_pict=True)
    blocks = load_docx(str(p))
    imgs = [b for b in blocks if b.kind == "image"]
    assert len(imgs) == 1
    import base64
    assert base64.b64decode(imgs[0].images[0]) == _PNG_1PX


# ---------------- parse：图片与题目关联 ----------------

import base64

from quizforge.model import Block
from quizforge.parse import parse_blocks

B64_PNG = base64.b64encode(_PNG_1PX).decode()


def test_image_between_stem_and_options_attaches():
    """题干与选项之间的图挂到该题（thermal 题 11 形态）。"""
    blocks = [
        Block("docx", "Heading 1", "单选题", "heading"),
        Block("docx", "Normal", "11、图示为静电除尘的原理示意图，尘埃受力方向为", "para"),
        Block("docx", "Normal", "[图片]", "image", images=[B64_PNG]),
        Block("docx", "Normal", "A、与电场方向相同", "para"),
        Block("docx", "Normal", "B、与电场方向相反", "para"),
        Block("docx", "Normal", "C、无关", "para"),
        Block("docx", "Normal", "D、不确定", "para"),
    ]
    quiz = parse_blocks(blocks, source_file="x.docx")
    q = quiz.questions[0]
    assert q.images == [B64_PNG], "题干后的图应挂到该题"
    assert "image" not in quiz.ignored_summary


def test_image_between_questions_attaches_to_next():
    """题与题之间的图挂到后一题（thermal 题 12 形态）。"""
    def opts(num, stem):
        return [
            Block("docx", "Normal", f"{num}、{stem}", "para"),
            Block("docx", "Normal", "A、甲", "para"),
            Block("docx", "Normal", "B、乙", "para"),
            Block("docx", "Normal", "C、丙", "para"),
            Block("docx", "Normal", "D、丁", "para"),
        ]

    blocks = [Block("docx", "Heading 1", "单选题", "heading")]
    blocks += opts(1, "第一题题干")
    blocks.append(Block("docx", "Normal", "[图片]", "image", images=[B64_PNG]))
    blocks += opts(2, "第二题题干")
    quiz = parse_blocks(blocks, source_file="x.docx")
    assert quiz.questions[0].images == []
    assert quiz.questions[1].images == [B64_PNG], "题间图应挂到后一题"


def test_knowledge_image_without_question_ignored():
    """不成题的图片（知识区）计入 ignored.image。"""
    blocks = [
        Block("docx", "Normal", "引言知识段落文字", "para"),
        Block("docx", "Normal", "[图片]", "image", images=[B64_PNG]),
        Block("docx", "Normal", "更多知识文字", "para"),
    ]
    quiz = parse_blocks(blocks, source_file="x.docx")
    assert not quiz.questions
    assert quiz.ignored_summary.get("image", 0) >= 1


def test_blank_in_choice_section_rescued_with_image():
    """单选节里混的填空题（无选项、空位在句中）→ 成存疑 blank 题并带配图（thermal 题 12 形态）。"""
    blocks = [
        Block("docx", "Heading 1", "单选题", "heading"),
        Block("docx", "Normal", "11、图示为静电除尘的原理示意图", "para"),
        Block("docx", "Normal", "A、甲", "para"),
        Block("docx", "Normal", "B、乙", "para"),
        Block("docx", "Normal", "C、丙", "para"),
        Block("docx", "Normal", "D、丁", "para"),
        Block("docx", "Normal", "[图片]", "image", images=[B64_PNG]),
        Block("docx", "Normal", "12、图示为电除尘气体电离和导电的伏安曲线图，若提高除尘效果，延长（  ）段的时间。", "para"),
    ]
    quiz = parse_blocks(blocks, source_file="x.docx")
    blanks = [q for q in quiz.questions if q.type == "blank"]
    assert len(blanks) == 1, "残缺填空应被救回为存疑 blank"
    b = blanks[0]
    assert "伏安曲线图" in b.stem
    assert b.images == [B64_PNG], "残缺填空应带上题间配图"
    assert b.low_confidence


def test_choice_stem_fragment_not_rescued():
    """句尾空括号的选择题残干（选项丢失）不救回，避免误产无意义填空。"""
    blocks = [
        Block("docx", "Heading 1", "单选题", "heading"),
        Block("docx", "Normal", "2、制冷机工作的共同特点是（  ）", "para"),
        Block("docx", "Normal", "能自动将能量由低温热源传送到高温热源；", "para"),
    ]
    quiz = parse_blocks(blocks, source_file="x.docx")
    assert not [q for q in quiz.questions if q.type == "blank"], "残干不应救回为填空"
